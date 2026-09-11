"""Test suite for ai_agent_core (run: manage.py test ai_agent_core)."""
from django.contrib.auth.models import Group, User
from django.core.exceptions import ValidationError
from django.test import RequestFactory, TestCase, override_settings

from .middleware import JupyterHubProxyMiddleware
from .models import TableAccessAudit, TableAccessPolicy
from .utils import (guarded_aggregate_data, guarded_describe_model,
                    guarded_fetch_data, guarded_list_models)

FIELDS = ["id", "username", "email", "is_active"]


class AccessControlTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice")
        cls.alice.groups.add(cls.hr)
        cls.bob = User.objects.create_user("bob")
        cls.root = User.objects.create_superuser("root", "r@x.com", "pw")
        cls.policy = TableAccessPolicy.objects.create(
            app_label="auth", model_name="User",
            access_level="read", allowed_fields=FIELDS, priority=10)
        cls.policy.groups.add(cls.hr)

    def test_group_member_can_fetch_with_narrowed_fields(self):
        r = guarded_fetch_data(self.alice, "auth.User")
        self.assertTrue(r["ok"])
        self.assertEqual(set(r["data"]["rows"][0]), set(FIELDS))

    def test_non_member_is_denied_and_audited(self):
        r = guarded_fetch_data(self.bob, "auth.User")
        self.assertFalse(r["ok"])
        self.assertEqual(r["error"]["code"], "PERMISSION_DENIED")
        self.assertTrue(TableAccessAudit.objects.filter(
            username="bob", was_allowed=False).exists())

    def test_filter_on_policy_hidden_field_is_blocked(self):
        r = guarded_fetch_data(self.alice, "auth.User",
                               filters={"last_login__isnull": True})
        self.assertFalse(r["ok"])
        self.assertIn("hidden", r["error"]["message"])

    def test_superuser_bypass(self):
        r = guarded_fetch_data(self.root, "auth.User", limit=2)
        self.assertTrue(r["ok"])

    def test_aggregate_and_discovery(self):
        self.assertEqual(guarded_aggregate_data(
            self.alice, "auth.User", "count", "id")["data"], 3)
        self.assertEqual(list(guarded_list_models(self.alice)["data"]),
                         ["auth.User"])
        desc = guarded_describe_model(self.alice, "auth.User")
        self.assertTrue(all(f["name"] in FIELDS
                            for f in desc["data"]["fields"]))

    def test_high_priority_deny_overrides_allow(self):
        TableAccessPolicy.objects.create(
            app_label="auth", model_name="User", access_level="deny",
            priority=99, applies_to_all_authenticated=True)
        self.assertFalse(guarded_fetch_data(self.alice, "auth.User")["ok"])

    def test_aggregate_only_blocks_row_fetch(self):
        self.policy.access_level = "aggregate_only"
        self.policy.save()
        self.assertFalse(guarded_fetch_data(self.alice, "auth.User")["ok"])
        self.assertTrue(guarded_aggregate_data(
            self.alice, "auth.User", "count", "id")["ok"])

    def test_policy_clean_rejects_non_allowlisted_table(self):
        with self.assertRaises(ValidationError):
            TableAccessPolicy(app_label="auth",
                              model_name="Nope").full_clean(exclude=["groups"])

    def test_policy_clean_rejects_non_allowlisted_field(self):
        with self.assertRaises(ValidationError):
            TableAccessPolicy(app_label="auth", model_name="User",
                              allowed_fields=["password"]
                              ).full_clean(exclude=["groups"])


class MiddlewareTests(TestCase):
    def _mw(self):
        return JupyterHubProxyMiddleware(lambda req: req)

    def test_prefix_header_is_stripped_into_script_name(self):
        req = RequestFactory().get(
            "/user/alice/aiagent/healthz/",
            HTTP_X_FORWARDED_PREFIX="/user/alice")
        req = self._mw()(req)
        self.assertEqual(req.META["SCRIPT_NAME"], "/user/alice")
        self.assertEqual(req.path_info, "/aiagent/healthz/")
        self.assertEqual(req.path, "/user/alice/aiagent/healthz/")

    def test_path_inference_without_header(self):
        req = RequestFactory().get("/user/bob/lab/")
        req = self._mw()(req)
        self.assertEqual(req.META["SCRIPT_NAME"], "/user/bob")

    @override_settings(AI_AGENT_TRUSTED_PROXY_COUNT=2)
    def test_client_ip_uses_trusted_hop_count(self):
        req = RequestFactory().get(
            "/", HTTP_X_FORWARDED_FOR="203.0.113.7, 10.0.0.5")
        req = self._mw()(req)
        self.assertEqual(req.META["AI_AGENT_CLIENT_IP"], "203.0.113.7")

    @override_settings(AI_AGENT_TRUSTED_PROXY_COUNT=2)
    def test_spoofed_extra_hops_do_not_fool_resolution(self):
        # A malicious client pre-set XFF; proxies appended the real chain.
        req = RequestFactory().get(
            "/", HTTP_X_FORWARDED_FOR="1.2.3.4, 203.0.113.7, 10.0.0.5")
        req = self._mw()(req)
        self.assertEqual(req.META["AI_AGENT_CLIENT_IP"], "203.0.113.7")


class DynamicRegistrationTests(TestCase):
    """Admin-registered tables + per-group field visibility."""

    @classmethod
    def setUpTestData(cls):
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice")
        cls.alice.groups.add(cls.hr)
        cls.bob = User.objects.create_user("bob")
        # Register a model that is NOT in the legacy allowlist.
        from .models import BotProfile, SearchableField, SearchableTable
        BotProfile.objects.create(name="Larry", greeting="Hello!",
                                  is_default=True)
        t = SearchableTable.objects.create(app_label="ai_agent_core",
                                           model_name="BotProfile")
        SearchableField.objects.create(table=t, field_name="id")
        SearchableField.objects.create(table=t, field_name="name")
        hidden = SearchableField.objects.create(table=t,
                                                field_name="greeting")
        hidden.groups.add(cls.hr)  # greeting only visible to HR
        from . import registry
        registry.mark_dirty()
        # Open the registered table to everyone via policy.
        registry.sync_registry(force=True)
        TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="BotProfile",
            access_level="read", applies_to_all_authenticated=True)

    def test_registered_table_is_searchable(self):
        r = guarded_fetch_data(self.alice, "ai_agent_core.BotProfile")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["rows"][0]["name"], "Larry")

    def test_group_restricted_field_visible_to_member(self):
        r = guarded_fetch_data(self.alice, "ai_agent_core.BotProfile")
        self.assertIn("greeting", r["data"]["rows"][0])

    def test_group_restricted_field_hidden_from_non_member(self):
        r = guarded_fetch_data(self.bob, "ai_agent_core.BotProfile")
        self.assertTrue(r["ok"], r)
        self.assertNotIn("greeting", r["data"]["rows"][0])
        self.assertIn("name", r["data"]["rows"][0])

    def test_non_member_cannot_filter_on_restricted_field(self):
        r = guarded_fetch_data(self.bob, "ai_agent_core.BotProfile",
                               filters={"greeting__icontains": "Hello"})
        self.assertFalse(r["ok"])

    def test_denied_field_names_cannot_be_registered(self):
        from .models import SearchableField, SearchableTable
        t = SearchableTable.objects.get()
        with self.assertRaises(ValidationError):
            SearchableField(table=t, field_name="password").full_clean(
                exclude=["groups"])

    def test_disabling_table_removes_access(self):
        from . import registry
        from .models import SearchableTable
        SearchableTable.objects.update(enabled=False)
        registry.mark_dirty()
        r = guarded_fetch_data(self.alice, "ai_agent_core.BotProfile")
        self.assertFalse(r["ok"])
        SearchableTable.objects.update(enabled=True)
        registry.mark_dirty()


class BotChatTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from .models import BotProfile, SearchableField, SearchableTable
        cls.bot = BotProfile.objects.create(
            name="Larry", greeting="Hello!", window_mode="popup",
            is_default=True)
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice", password="pw")
        cls.alice.groups.add(cls.hr)
        t = SearchableTable.objects.create(app_label="ai_agent_core",
                                           model_name="BotProfile")
        for f in ("id", "name", "greeting"):
            SearchableField.objects.create(table=t, field_name=f)
        from . import registry
        registry.mark_dirty()
        registry.sync_registry(force=True)
        p = TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="BotProfile",
            access_level="read")
        p.groups.add(cls.hr)

    def test_chat_endpoint_returns_reply_and_results(self):
        self.client.login(username="alice", password="pw")
        r = self.client.post("/chat/", {"message": "find Larry"},
                             content_type="application/json")
        self.assertEqual(r.status_code, 200)
        data = r.json()
        self.assertIn("Larry", str(data["results"]))
        self.assertTrue(data["reply"])
        self.assertIsInstance(data["conversation_id"], int)

    def test_chat_persists_conversation(self):
        from .models import BotChatMessage
        self.client.login(username="alice", password="pw")
        r1 = self.client.post("/chat/", {"message": "hello Larry"},
                              content_type="application/json").json()
        r2 = self.client.post(
            "/chat/", {"message": "more", "conversation_id":
                       r1["conversation_id"]},
            content_type="application/json").json()
        self.assertEqual(r1["conversation_id"], r2["conversation_id"])
        self.assertEqual(BotChatMessage.objects.filter(
            conversation_id=r1["conversation_id"]).count(), 4)

    def test_anonymous_chat_gets_graceful_no_access_reply(self):
        r = self.client.post("/chat/", {"message": "find Larry"},
                             content_type="application/json")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.json()["results"], [])

    def test_widget_demo_page_renders_launcher(self):
        r = self.client.get("/widget-demo/")
        self.assertContains(r, "aac-launcher")
        self.assertContains(r, "aac-popup")   # window_mode from BotProfile
        self.assertContains(r, "Larry")


class WidgetSvgTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from .models import BotProfile
        BotProfile.objects.create(name="Larry", greeting="Hello!",
                                  window_mode="right", is_default=True)

    def test_default_launcher_is_animated_svg(self):
        r = self.client.get("/widget-demo/")
        self.assertContains(r, "aac-robot")          # the SVG
        self.assertContains(r, "aac-antenna-light")  # pulsing antenna
        self.assertContains(r, "aac-eyelid")         # wink animation target
        self.assertContains(r, "aac-launcher aac-floating")  # fixed launcher
        self.assertContains(r, 'viewBox="2 2 80 90"')
        self.assertNotContains(r, "aac-launcher aac-inline")

    def test_inline_mode_fills_container(self):
        r = self.client.get("/widget-demo/?inline=1")
        self.assertContains(r, "aac-launcher aac-inline")
        self.assertNotContains(r, "aac-launcher aac-floating")

    def test_emoji_override_replaces_svg(self):
        from .models import BotProfile
        BotProfile.objects.update(avatar_emoji="\U0001F916")
        r = self.client.get("/widget-demo/")
        self.assertNotContains(r, '<svg class="aac-robot"')
        self.assertContains(r, "aac-emoji")

    def test_svg_click_opens_chat_wiring_present(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn('launcher.addEventListener("click"', content)
        self.assertIn("aac-panel", content)


class WidgetThemeTests(TestCase):
    @classmethod
    def setUpTestData(cls):
        from .models import BotProfile
        cls.bot = BotProfile.objects.create(name="Larry", greeting="Hi!",
                                            is_default=True)

    def test_auto_mode_renders_theme_machinery(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn('data-aac-mode="auto"', content)
        self.assertIn('data-aac-theme="dark"', content)     # dark palette CSS
        self.assertIn("prefers-color-scheme: dark", content)  # OS hook
        self.assertIn("data-bs-theme", content)             # Bootstrap hook
        self.assertIn("MutationObserver", content)          # live site toggles
        self.assertIn("window.aacSetTheme", content)        # manual hook

    def test_forced_dark_mode_is_rendered(self):
        from .models import BotProfile
        BotProfile.objects.update(theme_mode="dark")
        r = self.client.get("/widget-demo/")
        self.assertContains(r, 'data-aac-mode="dark"')

    def test_forced_light_mode_is_rendered(self):
        from .models import BotProfile
        BotProfile.objects.update(theme_mode="light")
        r = self.client.get("/widget-demo/")
        self.assertContains(r, 'data-aac-mode="light"')

    def test_dark_palette_uses_css_variables(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn("--aac-surface: #1f2937", content)
        self.assertIn("--aac-log-bg: #111827", content)
        # panel colors are variable-driven, not hardcoded
        self.assertIn("background: var(--aac-surface)", content)


class ChapteredReportTests(TestCase):
    """Big reports are written bit by bit with progress notes."""

    @classmethod
    def setUpTestData(cls):
        from .models import BotProfile, SearchableField, SearchableTable
        cls.bot = BotProfile.objects.create(name="Larry", greeting="Hi!",
                                            is_default=True)
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice", password="pw")
        cls.alice.groups.add(cls.hr)
        cls.mallory = User.objects.create_user("mallory", password="pw")
        t = SearchableTable.objects.create(app_label="ai_agent_core",
                                           model_name="BotProfile")
        for f in ("id", "name", "greeting"):
            SearchableField.objects.create(table=t, field_name=f)
        from . import registry
        registry.mark_dirty()
        registry.sync_registry(force=True)
        p = TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="BotProfile",
            access_level="read")
        p.groups.add(cls.hr)

    def _start_report(self):
        self.client.login(username="alice", password="pw")
        r = self.client.post(
            "/chat/",
            {"message": "generate a big report on Larry with 4 chapters"},
            content_type="application/json")
        return r.json()

    def test_report_request_creates_job_not_inline_answer(self):
        data = self._start_report()
        self.assertIn("report_job", data)
        self.assertIn("bit by bit", data["reply"])

    def test_stepping_writes_one_chapter_at_a_time_with_notes(self):
        from .models import ReportJob
        job_id = self._start_report()["report_job"]
        url = f"/reports/{job_id}/step/"

        r = self.client.post(url).json()          # outline pass
        self.assertEqual(r["total_chapters"], 4)
        self.assertIn("Writing chapter 1/4", r["note"])
        self.assertFalse(r["done"])

        r = self.client.post(url).json()          # chapter 1
        self.assertEqual(r["chapters_done"], 1)
        self.assertIn("Finished chapter 1/4", r["note"])
        self.assertIn("Writing chapter 2/4", r["note"])

        for _ in range(3):                        # chapters 2..4
            r = self.client.post(url).json()
        self.assertTrue(r["done"])
        self.assertIn("Report complete", r["note"])
        self.assertIn("download_url", r)
        job = ReportJob.objects.get(pk=job_id)
        self.assertEqual(job.chapters.filter(status="done").count(), 4)

    def test_download_contains_every_chapter(self):
        job_id = self._start_report()["report_job"]
        url = f"/reports/{job_id}/step/"
        for _ in range(5):
            self.client.post(url)
        r = self.client.get(f"/reports/{job_id}/download/")
        self.assertEqual(r.status_code, 200)
        md = r.content.decode()
        for i in range(1, 5):
            self.assertIn(f"Chapter {i}:", md)
        self.assertIn("Larry", md)   # guarded search data made it in

    def test_download_before_finish_is_409(self):
        job_id = self._start_report()["report_job"]
        r = self.client.get(f"/reports/{job_id}/download/")
        self.assertEqual(r.status_code, 409)

    def test_other_user_cannot_step_or_download(self):
        job_id = self._start_report()["report_job"]
        self.client.logout()
        self.client.login(username="mallory", password="pw")
        self.assertEqual(
            self.client.post(f"/reports/{job_id}/step/").status_code, 404)
        self.assertEqual(
            self.client.get(f"/reports/{job_id}/download/").status_code, 404)

    def test_chapter_context_uses_summaries_not_full_text(self):
        """The context-overflow fix: later chapters receive only summaries."""
        from .reports import step
        from .models import ReportJob
        job_id = self._start_report()["report_job"]
        job = ReportJob.objects.get(pk=job_id)
        captured = {}

        def spy_writer(mode, user, j, context):
            if mode == "outline":
                return [{"title": "A", "brief": ""},
                        {"title": "B", "brief": ""}]
            captured[context["chapter_index"]] = context
            return {"content": "word " * 500, "summary": "short summary"}

        from unittest import mock
        with mock.patch("ai_agent_core.reports.get_writer",
                        return_value=spy_writer):
            for _ in range(3):
                step(ReportJob.objects.get(pk=job_id), user=self.alice)

        ctx2 = captured[2]
        self.assertEqual(len(ctx2["previous_summaries"]), 1)
        self.assertEqual(ctx2["previous_summaries"][0]["summary"],
                         "short summary")
        blob = str(ctx2)
        self.assertNotIn("word word word word word", blob)  # no full text

    def test_widget_includes_report_polling(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn("runReportJob", content)
        self.assertIn("/reports/0/step/", content)


class ThemedDocumentTests(TestCase):
    """Word / PDF / Excel / PowerPoint output with a custom theme."""

    @classmethod
    def setUpTestData(cls):
        from .models import (BotProfile, DocumentTheme, SearchableField,
                             SearchableTable)
        cls.bot = BotProfile.objects.create(name="Larry", greeting="Hi!",
                                            is_default=True)
        cls.theme = DocumentTheme.objects.create(
            name="Acme Corporate", primary_color="#123456",
            secondary_color="#654321", accent_color="#FF8800",
            heading_font="Georgia", body_font="Verdana",
            footer_text="Acme Corp — Confidential", is_default=True)
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice", password="pw")
        cls.alice.groups.add(cls.hr)
        t = SearchableTable.objects.create(app_label="ai_agent_core",
                                           model_name="BotProfile")
        for f in ("id", "name", "greeting"):
            SearchableField.objects.create(table=t, field_name=f)
        from . import registry
        registry.mark_dirty()
        registry.sync_registry(force=True)
        p = TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="BotProfile",
            access_level="read")
        p.groups.add(cls.hr)

    def _chat(self, fmt):
        self.client.login(username="alice", password="pw")
        return self.client.post(
            "/chat/", {"message": "find Larry", "output": fmt},
            content_type="application/json").json()

    def _download(self, data):
        url = data["document"]["download_url"]
        path = url.split("//", 1)[-1].split("/", 1)[1]
        return self.client.get("/" + path)

    def test_widget_dropdown_has_powerpoint(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn('id="aac-output"', content)
        for v in ("chat", "word", "pdf", "excel", "ppt"):
            self.assertIn(f'value="{v}"', content)
        self.assertIn("PowerPoint", content)

    def test_word_document_generated_and_valid(self):
        import io
        from docx import Document
        data = self._chat("word")
        self.assertEqual(data["document"]["format"], "word")
        resp = self._download(data)
        d = Document(io.BytesIO(resp.content))
        text = "\n".join(p.text for p in d.paragraphs)
        self.assertIn("find Larry", text)

    def test_pdf_document_generated(self):
        data = self._chat("pdf")
        resp = self._download(data)
        self.assertTrue(resp.content.startswith(b"%PDF"))

    def test_excel_document_generated_with_data_sheet(self):
        import io
        from openpyxl import load_workbook
        data = self._chat("excel")
        wb = load_workbook(io.BytesIO(self._download(data).content))
        joined = " ".join(str(c.value) for ws in wb.worksheets
                          for row in ws.iter_rows() for c in row)
        self.assertIn("Larry", joined)

    def test_powerpoint_generated_with_custom_theme(self):
        import io
        from pptx import Presentation
        from pptx.util import Pt
        data = self._chat("ppt")
        self.assertEqual(data["document"]["format"], "ppt")
        prs = Presentation(io.BytesIO(self._download(data).content))
        self.assertGreaterEqual(len(prs.slides), 2)   # title + content
        texts, fonts, colors = [], set(), set()
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for p in shape.text_frame.paragraphs:
                        for run in p.runs:
                            texts.append(run.text)
                            if run.font.name:
                                fonts.add(run.font.name)
                            if run.font.color and run.font.color.type is not None:
                                colors.add(str(run.font.color.rgb))
        blob = " ".join(texts)
        self.assertIn("find Larry", blob)                    # title slide
        self.assertIn("Acme Corp — Confidential", blob)      # themed footer
        self.assertIn("Georgia", fonts)                      # heading font
        self.assertIn("Verdana", fonts)                      # body font
        self.assertIn("123456", {c.upper() for c in colors} |
                                {c for c in colors})         # primary color

    def test_powerpoint_aliases(self):
        for alias in ("pptx", "powerpoint"):
            data = self._chat(alias)
            self.assertEqual(data["document"]["format"], "ppt", alias)

    def test_report_downloads_as_powerpoint(self):
        self.client.login(username="alice", password="pw")
        job_id = self.client.post(
            "/chat/", {"message": "generate a report on Larry with 3 chapters"},
            content_type="application/json").json()["report_job"]
        for _ in range(4):
            self.client.post(f"/reports/{job_id}/step/")
        r = self.client.get(f"/reports/{job_id}/download/?format=ppt")
        self.assertEqual(r.status_code, 200)
        import io
        from pptx import Presentation
        prs = Presentation(io.BytesIO(r.content))
        self.assertGreaterEqual(len(prs.slides), 4)  # title + 3 chapters
        r = self.client.get(f"/reports/{job_id}/download/?format=word")
        self.assertEqual(r.status_code, 200)

    def test_document_owner_scoping(self):
        data = self._chat("word")
        self.client.logout()
        User.objects.create_user("mallory", password="pw")
        self.client.login(username="mallory", password="pw")
        self.assertEqual(self._download(data).status_code, 404)

    def test_chat_format_returns_no_document(self):
        data = self._chat("chat")
        self.assertNotIn("document", data)


class AuditProvenanceTests(TestCase):
    """Collapsible chat audit trail + AI provenance stamped on documents."""

    @classmethod
    def setUpTestData(cls):
        from .models import (BotProfile, DocumentTheme, SearchableField,
                             SearchableTable)
        cls.bot = BotProfile.objects.create(name="Larry", greeting="Hi!",
                                            is_default=True)
        DocumentTheme.objects.create(name="Acme", is_default=True)
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice", password="pw")
        cls.alice.groups.add(cls.hr)
        t = SearchableTable.objects.create(app_label="ai_agent_core",
                                           model_name="BotProfile")
        for f in ("id", "name", "greeting"):
            SearchableField.objects.create(table=t, field_name=f)
        from . import registry
        registry.mark_dirty()
        registry.sync_registry(force=True)
        p = TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="BotProfile",
            access_level="read")
        p.groups.add(cls.hr)

    def _chat(self, message="find Larry", fmt=None):
        self.client.login(username="alice", password="pw")
        body = {"message": message}
        if fmt:
            body["output"] = fmt
        return self.client.post("/chat/", body,
                                content_type="application/json").json()

    def test_chat_response_carries_sources_and_model(self):
        data = self._chat()
        self.assertTrue(data["sources"])
        hit = [s for s in data["sources"]
               if s["table"] == "ai_agent_core.BotProfile"]
        self.assertTrue(hit and hit[0]["hit"] and hit[0]["rows"] >= 1)
        self.assertIn("mode", data["model"])
        self.assertEqual(data["model"]["model"], "gpt-4o-mini")  # from profile

    def test_widget_renders_collapsible_audit_details(self):
        r = self.client.get("/widget-demo/")
        content = r.content.decode()
        self.assertIn("aac-sources", content)
        self.assertIn("<summary>", content.replace("&#128269; ", "")
                      if False else content)  # summary element present
        self.assertIn("addSources", content)
        self.assertIn("audit trail", content)
        self.assertIn("Generated by AI", content)

    def test_word_document_contains_audit_and_model(self):
        import io
        from docx import Document
        data = self._chat(fmt="word")
        path = data["document"]["download_url"].split("//", 1)[-1].split("/", 1)[1]
        d = Document(io.BytesIO(self.client.get("/" + path).content))
        text = "\n".join(p.text for p in d.paragraphs)
        cells = " ".join(c.text for tb in d.tables for row in tb.rows
                         for c in row.cells)
        self.assertIn("Generated by AI", text)
        self.assertIn("Audit & Provenance", text)
        self.assertIn("gpt-4o-mini", text)
        self.assertIn("Requested by: alice", text)
        self.assertIn("ai_agent_core.BotProfile", cells)  # sources table

    def test_powerpoint_has_audit_slide(self):
        import io
        from pptx import Presentation
        data = self._chat(fmt="ppt")
        path = data["document"]["download_url"].split("//", 1)[-1].split("/", 1)[1]
        prs = Presentation(io.BytesIO(self.client.get("/" + path).content))
        blob = " ".join(run.text for s in prs.slides for sh in s.shapes
                        if sh.has_text_frame for p in sh.text_frame.paragraphs
                        for run in p.runs)
        self.assertIn("Audit & Provenance", blob)
        self.assertIn("Generated by AI", blob)
        self.assertIn("gpt-4o-mini", blob)

    def test_report_markdown_and_pptx_have_audit_trail(self):
        self.client.login(username="alice", password="pw")
        job_id = self.client.post(
            "/chat/", {"message": "generate a report on Larry with 3 chapters"},
            content_type="application/json").json()["report_job"]
        for _ in range(4):
            self.client.post(f"/reports/{job_id}/step/")
        md = self.client.get(f"/reports/{job_id}/download/").content.decode()
        self.assertIn("## Audit & Provenance", md)
        self.assertIn("generated by AI", md)
        self.assertIn("Requested by: alice", md)
        self.assertIn("ai_agent_core.BotProfile", md)  # audited reads
        self.assertIn("allowed", md)
        import io
        from pptx import Presentation
        r = self.client.get(f"/reports/{job_id}/download/?format=ppt")
        prs = Presentation(io.BytesIO(r.content))
        blob = " ".join(run.text for s in prs.slides for sh in s.shapes
                        if sh.has_text_frame for p in sh.text_frame.paragraphs
                        for run in p.runs)
        self.assertIn("Audit & Provenance", blob)


class CalculationAuditTests(TestCase):
    """Aggregations/calculations land in the audit log with their results."""

    @classmethod
    def setUpTestData(cls):
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice")
        cls.alice.groups.add(cls.hr)
        p = TableAccessPolicy.objects.create(
            app_label="auth", model_name="User", access_level="read")
        p.groups.add(cls.hr)

    def test_aggregate_formula_and_result_are_logged(self):
        r = guarded_aggregate_data(self.alice, "auth.User", "count", "id",
                                   filters={"is_active": True})
        self.assertTrue(r["ok"])
        audit = TableAccessAudit.objects.filter(action="aggregate").latest(
            "created_at")
        self.assertTrue(audit.was_allowed)
        self.assertEqual(audit.query["func"], "count")     # the formula
        self.assertEqual(audit.query["field"], "id")
        self.assertEqual(audit.query["filters"], {"is_active": True})
        self.assertTrue(audit.query["ok"])
        self.assertEqual(audit.query["result"], r["data"])  # the outcome

    def test_grouped_aggregate_result_logged(self):
        r = guarded_aggregate_data(self.alice, "auth.User", "count", "id",
                                   group_by="is_active")
        self.assertTrue(r["ok"])
        audit = TableAccessAudit.objects.filter(action="aggregate").latest(
            "created_at")
        self.assertEqual(audit.query["group_by"], "is_active")
        self.assertEqual(audit.query["result"], r["data"])
        self.assertEqual(audit.row_count, len(r["data"]))

    def test_denied_aggregate_still_logged_with_formula(self):
        bob = User.objects.create_user("bob")
        r = guarded_aggregate_data(bob, "auth.User", "sum", "id")
        self.assertFalse(r["ok"])
        audit = TableAccessAudit.objects.filter(
            username="bob", action="aggregate").latest("created_at")
        self.assertFalse(audit.was_allowed)
        self.assertEqual(audit.query["func"], "sum")   # attempt is on record


class ComputeFormulaTests(TestCase):
    """Audited in-database formulas, incl. cross-table."""

    @classmethod
    def setUpTestData(cls):
        from .models import ReportChapter, ReportJob, SearchableField, SearchableTable
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice")
        cls.alice.groups.add(cls.hr)
        cls.bob = User.objects.create_user("bob")
        # Register two related tables (chapter -> job)
        t1 = SearchableTable.objects.create(app_label="ai_agent_core",
                                            model_name="ReportChapter")
        for f in ("id", "index", "word_count", "job"):
            SearchableField.objects.create(table=t1, field_name=f)
        t2 = SearchableTable.objects.create(app_label="ai_agent_core",
                                            model_name="ReportJob")
        for f in ("id", "total_chapters", "chapters_done"):
            SearchableField.objects.create(table=t2, field_name=f)
        from . import registry
        registry.mark_dirty()
        registry.sync_registry(force=True)
        for model in ("ReportChapter", "ReportJob"):
            p = TableAccessPolicy.objects.create(
                app_label="ai_agent_core", model_name=model,
                access_level="read")
            p.groups.add(cls.hr)
        # bob: chapters yes (but word_count hidden), jobs NO
        TableAccessPolicy.objects.create(
            app_label="ai_agent_core", model_name="ReportChapter",
            access_level="read", allowed_fields=["id", "index", "job"],
            applies_to_all_authenticated=True, priority=-1)
        # alice may also count users (for scalars)
        pu = TableAccessPolicy.objects.create(
            app_label="auth", model_name="User", access_level="read")
        pu.groups.add(cls.hr)
        cls.job = ReportJob.objects.create(request_text="x", title="T",
                                           total_chapters=4)
        for idx, wc in ((1, 100), (2, 200), (4, 0)):
            ReportChapter.objects.create(job=cls.job, index=idx,
                                         title=f"c{idx}", word_count=wc)

    def _c(self, user, **kw):
        from .compute import guarded_compute_data
        return guarded_compute_data(user, "ai_agent_core.ReportChapter", **kw)

    def test_row_formula_a_div_b_times_a(self):
        r = self._c(self.alice, expression="word_count / index * word_count",
                    order_by="index")
        self.assertTrue(r["ok"], r)
        by_idx = {row["index"]: row["result"] for row in r["data"]["rows"]}
        self.assertEqual(by_idx[1], 100 / 1 * 100)
        self.assertEqual(by_idx[2], 200 / 2 * 200)
        self.assertEqual(by_idx[4], 0.0)

    def test_division_by_zero_yields_null_not_error(self):
        r = self._c(self.alice, expression="index / word_count")
        self.assertTrue(r["ok"], r)
        by_idx = {row["index"]: row["result"] for row in r["data"]["rows"]}
        self.assertIsNone(by_idx[4])          # 4 / 0 -> null
        self.assertEqual(by_idx[1], 0.01)

    def test_aggregate_of_formula(self):
        r = self._c(self.alice, expression="word_count * 2", aggregate="sum")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["result"], 600.0)

    def test_grouped_aggregate(self):
        r = self._c(self.alice, expression="word_count", aggregate="sum",
                    group_by="job")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["result"][0]["result"], 300.0)

    def test_cross_table_via_relation_path(self):
        r = self._c(self.alice, expression="word_count / job__total_chapters",
                    filters={"index": 1})
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["rows"][0]["result"], 25.0)   # 100 / 4

    def test_cross_table_denied_when_related_table_not_granted(self):
        r = self._c(self.bob, expression="index / job__total_chapters")
        self.assertFalse(r["ok"])
        self.assertEqual(r["error"]["code"], "PERMISSION_DENIED")
        self.assertIn("ReportJob", r["error"]["message"])

    def test_cross_table_via_scalars_from_unrelated_table(self):
        r = self._c(self.alice, expression="word_count / total_users",
                    filters={"index": 2},
                    scalars={"total_users": {"model": "auth.User",
                                             "func": "count", "field": "id"}})
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["rows"][0]["result"], 200 / 2)  # 2 users
        self.assertEqual(r["data"]["scalars"][0]["value"], 2.0)

    def test_hidden_field_cannot_be_used_in_formula(self):
        r = self._c(self.bob, expression="word_count * 2")
        self.assertFalse(r["ok"])
        self.assertEqual(r["error"]["code"], "PERMISSION_DENIED")

    def test_unsafe_syntax_rejected(self):
        for bad in ("__import__('os').system('x')", "word_count ** 2",
                    "index.__class__", "max(index, 1)", "index > 1"):
            r = self._c(self.alice, expression=bad)
            self.assertFalse(r["ok"], bad)
            self.assertEqual(r["error"]["code"], "VALIDATION_ERROR", bad)

    def test_formula_and_result_are_audited(self):
        self._c(self.alice, expression="word_count / index", aggregate="avg")
        a = TableAccessAudit.objects.filter(action="compute").latest("created_at")
        self.assertTrue(a.was_allowed)
        self.assertEqual(a.query["expression"], "word_count / index")
        self.assertEqual(a.query["aggregate"], "avg")
        self.assertEqual(a.query["fields"], ["index", "word_count"])
        self.assertTrue(a.query["ok"])
        self.assertIn("result", a.query)
        # denied attempt also on record with the attempted formula
        self._c(self.bob, expression="word_count * 2")
        d = TableAccessAudit.objects.filter(action="compute",
                                            username="bob").latest("created_at")
        self.assertFalse(d.was_allowed)
        self.assertEqual(d.query["expression"], "word_count * 2")


class RouterIntegrationTests(TestCase):
    """Your own LLM router can use the audited tools (no Larry required)."""

    @classmethod
    def setUpTestData(cls):
        cls.hr = Group.objects.create(name="HR")
        cls.alice = User.objects.create_user("alice")
        cls.alice.groups.add(cls.hr)
        for model in ("User", "Group"):
            p = TableAccessPolicy.objects.create(
                app_label="auth", model_name=model, access_level="read")
            p.groups.add(cls.hr)

    def test_schemas_cover_legacy_tools_plus_compute(self):
        from .integrations import TOOL_SCHEMAS
        names = {s["function"]["name"] for s in TOOL_SCHEMAS}
        self.assertEqual(names, {"list_models", "describe_model",
                                 "fetch_data", "aggregate_data",
                                 "fetch_related", "compute_data"})

    def test_execute_tool_openai_style(self):
        import json
        from .integrations import execute_tool
        out = json.loads(execute_tool(
            "aggregate_data",
            json.dumps({"model_path": "auth.User", "func": "count",
                        "field": "id"}), user=self.alice))
        self.assertTrue(out["ok"])
        self.assertEqual(out["data"], 1)
        out = json.loads(execute_tool(
            "compute_data", {"model_path": "auth.User",
                             "expression": "id * 10"}, user=self.alice))
        self.assertTrue(out["ok"], out)
        self.assertEqual(out["data"]["rows"][0]["result"], self.alice.pk * 10)
        bad = json.loads(execute_tool("drop_table", "{}", user=self.alice))
        self.assertEqual(bad["error"]["code"], "UNKNOWN_TOOL")

    def test_fetch_related_is_guarded_and_audited(self):
        from .utils import guarded_fetch_related
        r = guarded_fetch_related(self.alice, "auth.User", self.alice.pk,
                                  "groups")
        self.assertTrue(r["ok"], r)
        self.assertEqual(r["data"]["rows"][0]["name"], "HR")
        self.assertTrue(TableAccessAudit.objects.filter(
            action="related", was_allowed=True).exists())
        bob = User.objects.create_user("bob")
        self.assertFalse(guarded_fetch_related(
            bob, "auth.User", self.alice.pk, "groups")["ok"])

    def test_rag_context_for_external_router(self):
        from .integrations import rag_context
        ctx = rag_context(self.alice, "alice")
        self.assertIn("sources", ctx)
        self.assertIn("model", ctx)
        self.assertIn("compute_data", ctx["tools"])
